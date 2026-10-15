# -*- coding: utf-8 -*-
"""Kiểm thử 4 bảng kết quả, mô hình nến và tiêu chí chiến lược rủi ro cao."""
import numpy as np
import pandas as pd
import pytest

from bo_loc import bon_bang as kf
from bo_loc.phan_tich.ky_thuat import mo_hinh_nen
from bo_loc.tieu_chi import kiem_tra


def _nen(ds):
    """ds: [(open, high, low, close)] – 10 nến giảm dần phía trước để có 'xu hướng giảm'."""
    truoc = [(60 - i, 60.5 - i, 58.5 - i, 59 - i) for i in range(10)]
    rows = truoc + ds
    return pd.DataFrame(rows, columns=["open", "high", "low", "close"],
                        index=pd.bdate_range("2026-09-01", periods=len(rows)))


def test_mo_hinh_nen():
    assert mo_hinh_nen(_nen([(50, 50.2, 48, 49), (49, 52, 48.8, 51.8)])) == "Bullish Engulfing ↑"
    assert mo_hinh_nen(_nen([(50, 50.2, 49, 49.5), (49.0, 49.7, 46, 49.6)])) == "Hammer ↑"
    assert mo_hinh_nen(_nen([(50, 50.2, 49, 49.5), (49.4, 52, 49.35, 49.9)])) == "Inverted Hammer ↑"
    assert mo_hinh_nen(_nen([(50, 50.2, 49, 49.5), (49.5, 49.7, 46, 49.6)])) == "Dragonfly Doji ↑"
    assert mo_hinh_nen(_nen([(50, 50.2, 49, 49.5), (49.5, 50.5, 48.5, 49.52)])) == "Doji –"


def _m(ma, **k):
    m = {"Mã": ma, "gia": 50.0, "nganh": "", "nhom": "cang_bien", "klgd3t": 1e6, "gtgd20": 50.0, "beta": 1.3,
         "rsi": 35.0, "eps_dung": 5000.0, "pe": 10.0, "von_hoa": 10000.0, "Đạt lọc": "✔", "Điểm %": 70.0,
         "dt_4q_ty": 2000.0, "ln_4q_ty": 500.0, "ln_g_quy": 12.0, "bien_gop": 30.0, "bvps": 30000.0, "pb": 1.6,
         "roe": 15.0, "kl_phien": 9e5, "ma20_huong": "Giảm", "ma50_huong": "Giảm", "ma100_huong": "Tăng",
         "mo_hinh_nen": "", "bd_1n": -0.01, "bd_1t": -0.03, "bd_2t": -0.05, "bd_1th": -0.1, "bd_3th": -0.2,
         "bd_6th": 0.1, "bd_9th": 0.2, "bd_1nam": 0.3, "bd_ytd": 0.25}
    m.update(k)
    return m


def test_lap_4_bang_dung_cot_va_xep_hang_rs():
    ds = [_m("AAA", bd_6th=0.5, bd_1nam=0.9), _m("BBB", bd_6th=-0.2, bd_1nam=-0.5), _m("CCC")]
    b = kf.lap_4_bang(ds, ["AAA", "CCC"])
    assert list(b) == ["TỔNG QUAN", "TÀI CHÍNH", "KỸ THUẬT", "BIẾN ĐỘNG", "PHÂN TÍCH TÀI CHÍNH"]
    assert list(b["BIẾN ĐỘNG"].columns) == ["MÃ", "GIÁ", "1 NGÀY", "1 TUẦN", "2 TUẦN", "1 THÁNG", "3 THÁNG",
                                            "6 THÁNG", "9 THÁNG", "1 NĂM", "TÍNH ĐẾN HIỆN TẠI"]
    assert list(b["KỸ THUẬT"]["MÃ"]) == ["AAA", "CCC"]                        # chỉ mã được chọn
    kt = b["KỸ THUẬT"].set_index("MÃ")
    assert kt.loc["AAA", "RS52W"] == 99 and kt.loc["AAA", "RS52W"] > kt.loc["CCC", "RS52W"]   # xếp hạng trên CẢ lần quét
    assert b["TỔNG QUAN"].loc[0, "Giá"] == 50000 and b["TỔNG QUAN"].loc[0, "NGÀNH"] == "Cảng biển"
    assert b["TÀI CHÍNH"].loc[0, "LỢI NHUẬN (TTM)"] == 500e9


def test_xuat_excel_tong_xanh_la(tmp_path):
    from openpyxl import load_workbook
    b = kf.lap_4_bang([_m("AAA"), _m("BBB")])
    tc = pd.DataFrame({"Loại": ["bat_buoc"], "Tiêu chí": ["Beta ≥ 1.2"]})
    f = kf.xuat_excel(str(tmp_path / "x.xlsx"), b, tc, "2 mã đạt", "Giá trị")
    wb = load_workbook(f)
    assert wb.sheetnames == ["TỔNG QUAN", "TÀI CHÍNH", "KỸ THUẬT", "BIẾN ĐỘNG", "TIÊU CHÍ"]   # chưa có FA → bỏ sheet
    ws = wb["BIẾN ĐỘNG"]
    assert ws["A1"].value == "BỘ LỌC CỔ PHIẾU – CHIẾN LƯỢC GIÁ TRỊ"
    assert ws["A4"].value == "MÃ" and ws["A4"].fill.fgColor.rgb == "FF1B5E20"           # tiêu đề xanh lá đậm
    assert ws["C5"].number_format == "0.####%" and ws["C5"].font.color.rgb == "FFC62828"    # âm → đỏ
    assert ws.freeze_panes == "B5" and wb["TIÊU CHÍ"]["A5"].value == "Bắt buộc"


def test_tieu_chi_rui_ro_cao():
    m = _m("AAA", rsi_vuot_30=True, cat_len_d=False, phan_ky_duong=False)
    assert kiem_tra("beta_min", m, 1.2) and not kiem_tra("beta_min", _m("X", beta=1.0), 1.2)
    assert kiem_tra("gia_trong", m, (10, 200)) and not kiem_tra("gia_trong", _m("X", gia=5), (10, 200))
    assert kiem_tra("eps_duong", m, None) and kiem_tra("eps_duong", _m("X", eps_dung=-10), None) is False
    assert kiem_tra("eps_duong", _m("X", eps_dung=np.nan), None) is None
    assert kiem_tra("tin_hieu_dao_chieu", m, None)
    m2 = _m("B", rsi_vuot_30=False, cat_len_d=False, phan_ky_duong=False, mo_hinh_nen="Hammer ↑")
    assert kiem_tra("tin_hieu_dao_chieu", m2, None)
    m3 = dict(m2, mo_hinh_nen="Shooting Star ↓")
    assert kiem_tra("tin_hieu_dao_chieu", m3, None) is False
    assert kiem_tra("klgd3t_min", m, 10_000)


def test_tieu_chi_gia_tri():
    m = _m("AAA", eps_dung=2500.0, pe=8.0, pe_tv_nhom=12.0)
    assert kiem_tra("eps_min", m, 2000) and kiem_tra("eps_min", _m("X", eps_dung=1500.0), 2000) is False
    assert kiem_tra("pe_duoi_tv_nganh", m, None)
    assert kiem_tra("pe_duoi_tv_nganh", _m("X", pe=15.0, pe_tv_nhom=12.0), None) is False
    assert kiem_tra("pe_duoi_tv_nganh", _m("X", pe_tv_nhom=np.nan), None) is None      # nhóm < 3 mã → thiếu dữ liệu


def _nen_tich_luy(so_phien=30, kl_cuoi=6e6, tang=0.06):
    """120 phiên tăng, rồi đi ngang 'so_phien' phiên (±3%), phiên cuối bật mạnh."""
    rng = np.random.default_rng(1)
    c = list(np.linspace(20, 40, 120)) + list(50 * (1 + 0.03 * np.sin(np.arange(so_phien))))
    c.append(c[-1] * (1 + tang) if tang else c[-1])
    c = np.array(c)
    df = pd.DataFrame({"open": c * 0.995, "high": c * 1.005, "low": c * 0.99, "close": c,
                       "volume": 1e6 * (1 + 0.1 * rng.random(len(c)))}, index=pd.bdate_range("2026-01-01", periods=len(c)))
    df.iloc[-1, df.columns.get_loc("volume")] = kl_cuoi
    df.iloc[-1, df.columns.get_loc("open")] = c[-2]
    return df


def test_tich_luy_va_vuot():
    from bo_loc.phan_tich.ky_thuat import so_phien_tich_luy, tich_luy_va_vuot
    d = _nen_tich_luy()
    assert 25 <= so_phien_tich_luy(d) <= 40
    kq = tich_luy_va_vuot(d, 20, 5, 5)
    assert kq["dat"] and kq["ty_le_kl"] > 5 and kq["tang_pct"] > 5
    assert not tich_luy_va_vuot(d, 40, 5, 5)["dat"]                         # nền chưa đủ 40 phiên
    assert not tich_luy_va_vuot(_nen_tich_luy(kl_cuoi=3e6), 20, 5, 5)["dat"]   # KL chỉ ×3
    assert not tich_luy_va_vuot(_nen_tich_luy(tang=0.03), 20, 5, 5)["dat"]     # giá chỉ +3%
    d2 = _nen_tich_luy(kl_cuoi=3.5e6)
    assert not tich_luy_va_vuot(d2, 20, 5, 5)["dat"] and tich_luy_va_vuot(d2, 20, 3, 5)["dat"]   # nới ngưỡng KL


def test_tieu_chi_tich_luy_vuot_qua_m_d():
    m = {"_d": _nen_tich_luy(), "tich_luy_phien": 30}
    assert kiem_tra("tich_luy_vuot", m, (20, 5, 5)) and kiem_tra("tich_luy_vuot", m, (40, 5, 5)) is False
    assert kiem_tra("dang_tich_luy", m, 20)
