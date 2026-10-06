# -*- coding: utf-8 -*-
"""
NHẬT KÝ KHUYẾN NGHỊ & CHẤM ĐIỂM ĐÚNG/SAI – MỘT bộ máy chấm dùng chung cho ptcp (Colab, danh-muc) và bot canh-bao-mua.

Mỗi lần ptcp.main() chạy (GHI_NHAT_KY = True):
  1. Chấm lại mọi khuyến nghị CŨ của mã đó bằng giá thực tế.
  2. Ghi khuyến nghị hôm nay nếu nó KHÁC lần ghi trước của mã (tránh trùng khi chạy nhiều lần / mẫu chồng nhau).
  3. Báo cáo có PHẦN K – lịch sử khuyến nghị của mã + độ chính xác.

Cách chấm (luật giao dịch VN như mô phỏng của ptcp):
  nhóm MUA        (MUA, MUA TỪNG PHẦN…)        → lệnh mua giả định ở giá MỞ CỬA phiên sau, cắt lỗ & mục tiêu NGẮN
                                                  HẠN của chính khuyến nghị; T+2; khoá trần không mua được, khoá sàn
                                                  không bán được; gap; cắt lỗ & mục tiêu cùng phiên → tính cắt lỗ.
                                                  ĐÚNG = chạm mục tiêu trước / hết hạn mà lãi sau phí > 0.
  nhóm CHỜ        (THEO DÕI, CHỜ…)              ┐ chấm bằng CHÍNH lệnh mua giả định đó nhưng ĐẢO kết quả:
  nhóm ĐỨNG NGOÀI (CHƯA MUA, KHÔNG MUA MỚI)     ┘ lệnh đó lỗ → khuyên đứng ngoài ĐÚNG; lệnh đó lãi → SAI (bỏ lỡ).
  BAN (cảnh báo bán của bot)                    → ĐÚNG nếu sau KY_HAN_BAN phiên giá đóng cửa ≤ giá lúc báo.
  Trạng thái: ĐÚNG · SAI · ĐANG CHỜ (kèm lãi/lỗ tạm tính) · BỎ QUA (không vào được lệnh / thiếu mức giá).
  ket_qua_pct: lãi/lỗ % sau phí của lệnh mua (thật hoặc giả định); với BAN = % đổi giá SAU khi báo (âm = bán đúng).

File (cfg.FILE_NHAT_KY = None → tự chọn):
  Colab đã mount Drive → /content/drive/MyDrive/ptcp/nhat_ky_ptcp.csv (không mất khi Colab tắt)
  không thì            → ./nhat_ky_ptcp.csv
"""
import os

import numpy as np
import pandas as pd

from . import cau_hinh as cfg

COT = ["id", "thoi_diem", "ngay", "ma", "loai", "khuyen_nghi", "nhom", "san", "gia", "cat_lo", "muc_tieu", "ev",
       "rr", "ky_han", "ket_qua", "giai_thich", "ngay_ket_thuc", "gia_ket_thuc", "ket_qua_pct", "so_phien"]
COT_CHU = ("id", "thoi_diem", "ngay", "ma", "loai", "khuyen_nghi", "nhom", "san", "ket_qua", "giai_thich",
           "ngay_ket_thuc")
CHUA_XONG = ("", "ĐANG CHỜ")
DUNG, SAI, CHO, BO_QUA = "ĐÚNG", "SAI", "ĐANG CHỜ", "BỎ QUA"
THU_MUC_DRIVE = "/content/drive/MyDrive"


def duong_dan():
    if cfg.FILE_NHAT_KY:
        return cfg.FILE_NHAT_KY
    if os.path.isdir(THU_MUC_DRIVE):
        return os.path.join(THU_MUC_DRIVE, "ptcp", "nhat_ky_ptcp.csv")
    return "nhat_ky_ptcp.csv"


def chi_phi():
    """% chi phí khứ hồi: phí + thuế + trượt giá 2 chiều (giống mô phỏng)."""
    return cfg.PHI_GD_KHU_HOI + 2 * cfg.TRUOT_GIA_PCT


# ------------------------------------------------------------------ đọc / ghi
def doc(path=None):
    path = path or duong_dan()
    if not os.path.exists(path):
        return pd.DataFrame(columns=COT)
    df = pd.read_csv(path, dtype={c: str for c in COT_CHU}, encoding="utf-8")
    for c in COT:
        if c not in df:
            df[c] = np.nan
    df[list(COT_CHU)] = df[list(COT_CHU)].fillna("")
    return df[COT]


def luu(df, path=None):
    path = path or duong_dan()
    d = os.path.dirname(path)
    if d:
        os.makedirs(d, exist_ok=True)
    df[COT].to_csv(path, index=False, encoding="utf-8")


def _so(x, n=2):
    try:
        x = float(x)
        return round(x, n) if x == x else np.nan
    except (TypeError, ValueError):
        return np.nan


def them_dong(dong, path=None):
    """Thêm 1 tín hiệu (bỏ qua nếu id đã có). Trả True nếu đã ghi."""
    df = doc(path)
    if dong["id"] in set(df["id"]):
        return False
    dong = {c: dong.get(c, "" if c in COT_CHU else np.nan) for c in COT}
    luu(pd.concat([df, pd.DataFrame([dong])], ignore_index=True) if len(df) else pd.DataFrame([dong]), path)
    return True


def nhom_khuyen_nghi(kn):
    kn = str(kn or "").upper()
    if kn.startswith("MUA"):
        return "MUA"
    if kn.startswith("CHƯA MUA") or kn.startswith("KHÔNG MUA"):
        return "ĐỨNG NGOÀI"
    return "CHỜ"


def ghi_khuyen_nghi(symbol, ngay, gia, khuyen_nghi, cat_lo, muc_tieu, ky_han, san="HOSE", ev=None, rr=None,
                    loai="PTCP", path=None, thoi_diem=None):
    """Ghi khuyến nghị nếu KHÁC lần ghi trước của mã (cùng loại). Trả True nếu đã ghi."""
    cu = doc(path)
    cu = cu[(cu["ma"] == symbol) & (cu["loai"] == loai)]
    if len(cu) and str(cu["khuyen_nghi"].iloc[-1]) == str(khuyen_nghi):
        return False
    ngay = f"{pd.Timestamp(ngay):%Y-%m-%d}"
    return them_dong({"id": f"{loai}|{symbol}|{ngay}", "ngay": ngay, "ma": symbol, "loai": loai,
                      "thoi_diem": thoi_diem or f"{pd.Timestamp.now():%Y-%m-%d %H:%M}",
                      "khuyen_nghi": str(khuyen_nghi), "nhom": nhom_khuyen_nghi(khuyen_nghi), "san": san or "HOSE",
                      "gia": _so(gia), "cat_lo": _so(cat_lo), "muc_tieu": _so(muc_tieu), "ev": _so(ev),
                      "rr": _so(rr), "ky_han": int(ky_han)}, path)


# ------------------------------------------------------------------ chấm điểm
def _khoa(dn, san=None):
    """Phiên khoá trần (không mua được) / khoá sàn (không bán được) – cùng định nghĩa với mô phỏng."""
    thr = ((cfg.BIEN_DO_SAN.get(str(san).upper(), 7.0) - 0.3) if san else cfg.NGUONG_TRAN_SAN) / 100
    c, l = dn["close"].values, dn["low"].values
    ct = np.r_[np.nan, c[:-1]]
    with np.errstate(invalid="ignore"):
        tran = l >= ct * (1 + thr)
        san_k = (c <= ct * (1 - thr)) & (c <= l * 1.0005)
    return np.nan_to_num(tran).astype(bool), np.nan_to_num(san_k).astype(bool)


def _kq(ket_qua, giai_thich, ngay_kt=None, gia_kt=np.nan, pct=np.nan, so_phien=np.nan):
    return {"ket_qua": ket_qua, "giai_thich": giai_thich,
            "ngay_ket_thuc": f"{pd.Timestamp(ngay_kt):%Y-%m-%d}" if ngay_kt is not None else "",
            "gia_ket_thuc": _so(gia_kt), "ket_qua_pct": _so(pct), "so_phien": so_phien}


def cham_lenh_mua(dn, ngay, gia, cat_lo, muc_tieu, ky_han, vao_mo_cua=True, san=None):
    """
    Mô phỏng 1 lệnh mua sau tín hiệu. vao_mo_cua=True: mua giá mở cửa phiên SAU `ngay` (khuyến nghị cuối ngày);
    False: mua ở `gia` trong phiên `ngay` (tín hiệu trong phiên), chỉ xét chạm giá từ phiên sau.
    """
    if not (cat_lo == cat_lo and muc_tieu == muc_tieu) or cat_lo >= gia:
        return _kq(BO_QUA, "thiếu cắt lỗ / mục tiêu")
    idx = dn.index.normalize()
    ngay = pd.Timestamp(ngay).normalize()
    o, h, l, c = (dn[k].values for k in ("open", "high", "low", "close"))
    tran, san_k = _khoa(dn, san)
    if vao_mo_cua:
        sau = np.flatnonzero(idx > ngay)
        if not len(sau):
            return _kq(CHO, "chưa có phiên vào lệnh")
        i0 = sau[0]
        if tran[i0]:
            return _kq(BO_QUA, f"phiên {idx[i0]:%d/%m} khoá trần – không mua được", idx[i0])
        vao, j0 = o[i0], i0
        if vao >= muc_tieu:
            return _kq(BO_QUA, f"mở cửa {vao:.2f} đã vượt mục tiêu", idx[i0])
    else:
        cung = np.flatnonzero(idx >= ngay)
        if not len(cung):
            return _kq(CHO, "chưa có dữ liệu phiên báo")
        i0, vao = cung[0], gia
        j0 = i0 + 1
    ks = None
    for j in range(j0, len(c)):
        ban_duoc = (j - i0) >= cfg.T_CONG and not san_k[j]
        if ks is None and l[j] <= cat_lo:
            ks = j
        if ks is not None:
            if ban_duoc:
                ra = min(cat_lo, o[j]) if j == ks else o[j]
                return _kq(SAI, "chạm cắt lỗ" + ("" if j == ks else " (bán trễ do T+2/khoá sàn)"),
                           idx[j], ra, (ra / vao - 1) * 100 - chi_phi(), j - i0)
            continue
        if ban_duoc and h[j] >= muc_tieu:
            ra = max(muc_tieu, o[j])
            return _kq(DUNG, "chạm mục tiêu", idx[j], ra, (ra / vao - 1) * 100 - chi_phi(), j - i0)
        if j - i0 >= ky_han and ban_duoc:
            pct = (c[j] / vao - 1) * 100 - chi_phi()
            return _kq(DUNG if pct > 0 else SAI, f"hết hạn {ky_han} phiên", idx[j], c[j], pct, j - i0)
    pct = (c[-1] / vao - 1) * 100 - chi_phi()
    return _kq(CHO, ("đã chạm cắt lỗ, chờ bán được" if ks is not None else
                     f"tạm tính {len(c) - 1 - i0}/{ky_han} phiên"), None, c[-1], pct, len(c) - 1 - i0)


def cham_ban(dn, ngay, gia, ky_han):
    idx = dn.index.normalize()
    cung = np.flatnonzero(idx >= pd.Timestamp(ngay).normalize())
    if not len(cung):
        return _kq(CHO, "chưa có dữ liệu")
    i0, c = cung[0], dn["close"].values
    j = i0 + int(ky_han)
    if j >= len(c):
        return _kq(CHO, f"tạm tính {len(c) - 1 - i0}/{ky_han} phiên", None, c[-1], (c[-1] / gia - 1) * 100,
                   len(c) - 1 - i0)
    pct = (c[j] / gia - 1) * 100
    return _kq(DUNG if pct <= 0 else SAI, ("giá giảm sau khi bán" if pct <= 0 else "giá tăng sau khi bán"),
               idx[j], c[j], pct, j - i0)


def cham_dong(r, dn):
    """Chấm 1 dòng nhật ký (Series) với giá ngày dn."""
    if r["loai"] == "BAN":
        return cham_ban(dn, r["ngay"], float(r["gia"]), int(r["ky_han"]))
    kq = cham_lenh_mua(dn, r["ngay"], float(r["gia"]), float(r["cat_lo"]), float(r["muc_tieu"]), int(r["ky_han"]),
                       vao_mo_cua=(r["loai"] != "MUA_NGAY"), san=r.get("san") or None)
    if r["nhom"] in ("CHỜ", "ĐỨNG NGOÀI") and kq["ket_qua"] == BO_QUA and "vượt mục tiêu" in kq["giai_thich"]:
        kq["ket_qua"], kq["giai_thich"] = SAI, "bỏ lỡ: phiên sau mở cửa đã vượt mục tiêu"
        return kq
    if r["nhom"] in ("CHỜ", "ĐỨNG NGOÀI") and kq["ket_qua"] in (DUNG, SAI):
        lai = kq["ket_qua_pct"] > 0                       # khuyên KHÔNG mua: đúng khi lệnh mua giả định lỗ
        kq["ket_qua"] = SAI if lai else DUNG
        kq["giai_thich"] = (f"bỏ lỡ: nếu mua thì {kq['giai_thich']}" if lai
                            else f"tránh được: nếu mua thì {kq['giai_thich']}")
    return kq


def cap_nhat(tai_gia, du_lieu_san=None, paths=None, chi_ma=None):
    """Chấm lại các dòng chưa có kết luận. tai_gia(ma) → giá ngày. chi_ma: chỉ chấm mã này. Trả số dòng vừa xong."""
    du_lieu_san, xong = dict(du_lieu_san or {}), 0
    for path in (paths or [duong_dan()]):
        df = doc(path)
        if not len(df):
            continue
        cho = df["ket_qua"].isin(CHUA_XONG) & ((df["ma"] == chi_ma) if chi_ma else True)
        if not cho.any():
            continue
        for i, r in df[cho].iterrows():
            if r["ma"] not in du_lieu_san:
                du_lieu_san[r["ma"]] = tai_gia(r["ma"]) if tai_gia else None
            dn = du_lieu_san[r["ma"]]
            if dn is None or not len(dn):
                continue
            kq = cham_dong(r, dn)
            for k, v in kq.items():
                df.at[i, k] = v
            xong += kq["ket_qua"] not in CHUA_XONG
        luu(df, path)
    return xong


# ------------------------------------------------------------------ thống kê
def _nhom_tk(df):
    return np.where(df["loai"] == "PTCP", "ptcp " + df["nhom"].astype(str),
                    np.where(df["loai"] == "BAN", "BÁN " + df["khuyen_nghi"].astype(str),
                             np.where(df["loai"] == "MUA_NGAY", "MUA NGAY", df["loai"].astype(str))))


def thong_ke(df):
    """Bảng độ chính xác theo loại / nhóm khuyến nghị."""
    if not len(df):
        return pd.DataFrame()
    df = df.copy()
    df["nhom_tk"] = _nhom_tk(df)
    rows = []
    for ten, g in df.groupby("nhom_tk", sort=False):
        xong = g[g["ket_qua"].isin([DUNG, SAI])]
        r = pd.to_numeric(xong["ket_qua_pct"], errors="coerce")
        rows.append({"Nhóm": ten, "Số tín hiệu": len(g), "Đã chấm": len(xong),
                     "Đúng": int((xong["ket_qua"] == DUNG).sum()), "Sai": int((xong["ket_qua"] == SAI).sum()),
                     "Đang chờ": int(g["ket_qua"].isin(CHUA_XONG).sum()), "Bỏ qua": int((g["ket_qua"] == BO_QUA).sum()),
                     "Tỷ lệ đúng %": (xong["ket_qua"] == DUNG).mean() * 100 if len(xong) else np.nan,
                     "TB kết quả %": r.mean() if len(r) else np.nan})
    return pd.DataFrame(rows)


def dong_thong_ke(df, tieu_de="📊 ĐỘ CHÍNH XÁC TÍN HIỆU"):
    tk = thong_ke(df)
    if not len(tk):
        return []
    d = [tieu_de]
    for _, r in tk.iterrows():
        if r["Đã chấm"]:
            nhan = "TB nếu mua" if ("CHỜ" in r["Nhóm"] or "ĐỨNG NGOÀI" in r["Nhóm"]) else \
                ("TB giá sau báo" if r["Nhóm"].startswith("BÁN") else "TB")
            d.append(f"  {r['Nhóm']}: đúng {r['Đúng']}/{r['Đã chấm']} ({r['Tỷ lệ đúng %']:.0f}%) | "
                     f"{nhan} {r['TB kết quả %']:+.1f}%" + (f" | chờ {r['Đang chờ']}" if r["Đang chờ"] else "")
                     + (" ⚠ ít mẫu" if r["Đã chấm"] < cfg.SO_MAU_TIN_CAY else ""))
        else:
            d.append(f"  {r['Nhóm']}: {r['Số tín hiệu']} tín hiệu, chưa có kết luận")
    return d


def trang_tri_excel(w, sheet_nhat_ky=None):
    """Tông xanh cho các sheet; tô màu dòng nhật ký theo kết quả."""
    from openpyxl.styles import Font, PatternFill
    for ws in w.book.worksheets:
        for o in ws[1]:
            o.font, o.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2E7D32")
        ws.freeze_panes = "A2"
        for cot in ws.columns:
            ws.column_dimensions[cot[0].column_letter].width = min(40, max(10, *(len(str(o.value or ""))
                                                                                for o in cot[:50])) + 2)
    if sheet_nhat_ky and sheet_nhat_ky in w.book.sheetnames:
        ws = w.book[sheet_nhat_ky]
        tieu_de = [o.value for o in ws[1]]
        if "ket_qua" in tieu_de:
            k = tieu_de.index("ket_qua")
            mau = {DUNG: "C8E6C9", SAI: "FFCDD2", CHO: "FFF9C4", BO_QUA: "EEEEEE"}
            for hang in ws.iter_rows(min_row=2):
                m = mau.get(hang[k].value)
                if m:
                    for o in hang:
                        o.fill = PatternFill("solid", fgColor=m)


# ------------------------------------------------------------------ tóm tắt theo mã (giá tới HIỆN TẠI)
def thu_muc_luu():
    """Thư mục lưu báo cáo: cạnh file nhật ký (Drive khi đã mount)."""
    return os.path.dirname(os.path.abspath(duong_dan()))


def lam_moi_gia(paths=None, tai_gia=None, im_lang=True):
    """
    Tải giá ngày MỚI NHẤT cho mọi mã trong nhật ký, chấm lại các dòng chưa kết luận (lãi/lỗ tạm tính cập nhật tới
    phiên gần nhất) và GHI LẠI file nhật ký (trên Drive nếu đã mount). Trả {mã: DataFrame giá ngày}.
    """
    paths = paths or [duong_dan()]
    nk = pd.concat([doc(p) for p in paths], ignore_index=True)
    if not len(nk):
        return {}
    if tai_gia is None:
        from .du_lieu import _tai_ngay

        def tai_gia(ma, start):
            return _tai_ngay(ma, start, "NGÀY", im_lang=im_lang)
    gia = {}
    for ma, g in nk.groupby("ma"):
        start = f"{pd.to_datetime(g['ngay']).min() - pd.Timedelta(days=20):%Y-%m-%d}"
        try:
            dn = tai_gia(ma, start)
            if dn is not None and len(dn):
                gia[ma] = dn
        except Exception as e:                                # 1 mã lỗi không làm hỏng cả bảng
            if not im_lang:
                print(f"  ⚠ {ma}: không tải được giá ({str(e)[:80]})")
    cap_nhat(None, gia, paths)
    return gia


def _danh_gia_tam(nhom, pct, nguong=1.0):
    """Đánh giá theo hướng giá TỪ LÚC khuyến nghị tới nay (chưa phải kết luận chấm)."""
    if pct != pct:
        return ""
    if abs(pct) < nguong:
        return "– đi ngang"
    if nhom == "MUA":
        return "✔ đúng hướng" if pct > 0 else "✘ ngược hướng"
    if nhom in ("ĐỨNG NGOÀI", "CHỜ"):
        return "✔ tránh được giảm" if pct < 0 else "✘ đang bỏ lỡ"
    return "✔ đúng hướng" if pct < 0 else "✘ ngược hướng"     # cảnh báo BÁN


def tom_tat_theo_ma(nk, gia=None):
    """
    nk: nhật ký; gia: {mã: giá ngày} (thiếu mã → dùng giá kết thúc / tạm tính trong nhật ký).
    Trả (bang_ma, bang_chi_tiet):
      bang_ma      : MỖI MÃ 1 dòng – khuyến nghị GẦN NHẤT, giá lúc đó → giá hiện tại, % tăng/giảm, đánh giá tạm,
                     % từ lần phân tích đầu tiên, số khuyến nghị & đúng/sai của mã.
      bang_chi_tiet: MỖI khuyến nghị 1 dòng – giá lúc khuyến nghị → hiện tại, kết quả chấm.
    """
    gia = gia or {}
    if not len(nk):
        return pd.DataFrame(), pd.DataFrame()
    d = nk.copy()
    d["ngay_dt"] = pd.to_datetime(d["ngay"])
    d["gia"] = pd.to_numeric(d["gia"], errors="coerce")

    def gia_nay(ma, r=None):
        dn = gia.get(ma)
        if dn is not None and len(dn):
            return float(dn["close"].iloc[-1]), dn.index[-1]
        if r is not None and r["gia_ket_thuc"] == r["gia_ket_thuc"]:
            return float(r["gia_ket_thuc"]), pd.NaT
        return np.nan, pd.NaT

    ct = []
    for _, r in d.sort_values(["ma", "ngay_dt"]).iterrows():
        g, ngay_g = gia_nay(r["ma"], r)
        pct = (g / r["gia"] - 1) * 100 if r["gia"] == r["gia"] and r["gia"] else np.nan
        ct.append({"Mã": r["ma"], "Ngày KN": r["ngay_dt"], "Khuyến nghị": r["khuyen_nghi"], "Nhóm": r["nhom"],
                   "Giá lúc KN": r["gia"], "Giá hiện tại": g, "Thay đổi %": pct,
                   "Đánh giá tạm": _danh_gia_tam(r["nhom"], pct), "Cắt lỗ": _so(r["cat_lo"]),
                   "Mục tiêu": _so(r["muc_tieu"]), "Kết quả chấm": r["ket_qua"] or CHO,
                   "Lệnh giả định %": _so(r["ket_qua_pct"]), "Ghi chú": r["giai_thich"], "_ngay_gia": ngay_g})
    ct = pd.DataFrame(ct)
    ma = []
    for m, g in ct.groupby("Mã", sort=False):
        cuoi, dau = g.iloc[-1], g.iloc[0]
        xong = g[g["Kết quả chấm"].isin([DUNG, SAI])]
        ma.append({"Mã": m, "Khuyến nghị gần nhất": cuoi["Khuyến nghị"], "Ngày KN": cuoi["Ngày KN"],
                   "Giá lúc KN": cuoi["Giá lúc KN"], "Giá hiện tại": cuoi["Giá hiện tại"],
                   "Thay đổi %": cuoi["Thay đổi %"], "Đánh giá tạm": cuoi["Đánh giá tạm"],
                   "Từ lần đầu %": (cuoi["Giá hiện tại"] / dau["Giá lúc KN"] - 1) * 100 if dau["Giá lúc KN"] else np.nan,
                   "Lần đầu": dau["Ngày KN"], "Số KN": len(g),
                   "Đúng/Sai": f"{int((xong['Kết quả chấm'] == DUNG).sum())}/{int((xong['Kết quả chấm'] == SAI).sum())}",
                   "Giá tới ngày": cuoi["_ngay_gia"]})
    bang_ma = pd.DataFrame(ma).sort_values("Thay đổi %", ascending=False, na_position="last").reset_index(drop=True)
    return bang_ma, ct.drop(columns=["_ngay_gia"]).sort_values(["Ngày KN", "Mã"], ascending=[False, True]) \
        .reset_index(drop=True)


def _to_mau_pct(v):
    if not isinstance(v, (int, float)) or v != v:
        return ""
    if v >= 1:
        return "color:#1B5E20;font-weight:bold;background-color:#E8F5E9"
    if v <= -1:
        return "color:#B71C1C;font-weight:bold;background-color:#FFEBEE"
    return "color:#555"


def _to_mau_danh_gia(v):
    v = str(v)
    return ("color:#1B5E20;font-weight:bold" if v.startswith("✔") else
            "color:#B71C1C;font-weight:bold" if v.startswith("✘") else "")


def kieu_bang(bang):
    """Styler tông xanh cho Colab: % xanh khi tăng, đỏ khi giảm; đánh giá ✔/✘ tô màu."""
    so = {c: "{:,.2f}" for c in ("Giá lúc KN", "Giá hiện tại", "Cắt lỗ", "Mục tiêu") if c in bang}
    so.update({c: "{:+.1f}%" for c in ("Thay đổi %", "Từ lần đầu %", "Lệnh giả định %") if c in bang})
    so.update({c: (lambda x: "" if pd.isna(x) else f"{x:%d/%m/%Y}") for c in ("Ngày KN", "Lần đầu", "Giá tới ngày")
               if c in bang})
    st = bang.style.format(so, na_rep="–").hide(axis="index")
    ap = getattr(st, "map", None) or st.applymap
    st = ap(_to_mau_pct, subset=[c for c in ("Thay đổi %", "Từ lần đầu %", "Lệnh giả định %") if c in bang])
    ap = getattr(st, "map", None) or st.applymap
    st = ap(_to_mau_danh_gia, subset=[c for c in ("Đánh giá tạm",) if c in bang])
    return st.set_table_styles([
        {"selector": "th", "props": "background-color:#2E7D32;color:white;font-weight:bold;text-align:center"},
        {"selector": "td", "props": "padding:4px 10px"}])


def xem(lam_moi=True, paths=None, xuat=True, chi_tiet=True, tai_gia=None):
    """
    BẢNG ĐIỀU KHIỂN NHẬT KÝ cho Colab – một lệnh:
      1. tải giá mới nhất mọi mã & chấm lại → ghi file nhật ký (Drive),
      2. hiện: tóm tắt THEO MÃ (tăng/giảm bao nhiêu từ lúc khuyến nghị tới nay) · độ chính xác theo nhóm ·
         chi tiết từng khuyến nghị,
      3. xuất danh_gia_khuyen_nghi.xlsx vào CÙNG thư mục với nhật ký (Drive).
    """
    paths = paths or [duong_dan()]
    gia = lam_moi_gia(paths, tai_gia) if lam_moi else {}
    nk = pd.concat([doc(p) for p in paths], ignore_index=True)
    bang_ma, ct = tom_tat_theo_ma(nk, gia)
    tk = thong_ke(nk)
    file_xlsx = xuat_excel(paths=paths, gia=gia) if xuat and len(nk) else None
    try:
        from IPython.display import HTML, display
        ngay = max((v.index[-1] for v in gia.values()), default=None)
        tang = int((bang_ma["Thay đổi %"] >= 1).sum()) if len(bang_ma) else 0
        giam = int((bang_ma["Thay đổi %"] <= -1).sum()) if len(bang_ma) else 0
        display(HTML(f"<h3 style='color:#1B5E20;margin:4px 0'>📒 NHẬT KÝ KHUYẾN NGHỊ – {len(bang_ma)} mã"
                     + (f" · giá tới {ngay:%d/%m/%Y}" if ngay is not None else "") + f" · ▲ {tang} tăng · ▼ {giam} "
                     f"giảm</h3><div style='color:#555'>Thay đổi % = giá hiện tại so với giá lúc khuyến nghị gần nhất."
                     f" Đánh giá tạm: MUA → giá tăng là đúng hướng; CHỜ / ĐỨNG NGOÀI → giá giảm là tránh được.</div>"))
        if len(bang_ma):
            display(kieu_bang(bang_ma))
        display(HTML("<h4 style='color:#1B5E20;margin:12px 0 4px'>Độ chính xác theo nhóm (đã chấm xong)</h4>"))
        display(tk)
        if chi_tiet and len(ct):
            display(HTML("<h4 style='color:#1B5E20;margin:12px 0 4px'>Chi tiết từng khuyến nghị</h4>"))
            display(kieu_bang(ct.drop(columns=["Nhóm"])))
        display(HTML(f"<div style='color:#2E7D32'>✔ Nhật ký: {paths[0]}"
                     + (f"<br>✔ Excel: {file_xlsx}" if file_xlsx else "")
                     + ("" if "/drive/" in paths[0] else "<br>⚠ Chưa mount Drive – file mất khi tắt Colab: "
                        "from google.colab import drive; drive.mount('/content/drive')") + "</div>"))
    except ImportError:                                        # ngoài notebook → in chữ
        print(bang_ma.to_string(index=False))
    return {"theo_ma": bang_ma, "chi_tiet": ct, "thong_ke": tk, "file_excel": file_xlsx, "path": paths[0]}


def xuat_excel(path_xlsx=None, paths=None, gia=None):
    """Excel tông xanh: Tóm tắt theo mã (tô xanh/đỏ) · Thống kê · Chi tiết · Nhật ký. Mặc định lưu cạnh nhật ký."""
    path_xlsx = path_xlsx or os.path.join(thu_muc_luu(), "danh_gia_khuyen_nghi.xlsx")
    nk = pd.concat([doc(p) for p in (paths or [duong_dan()])], ignore_index=True)
    bang_ma, ct = tom_tat_theo_ma(nk, gia)
    d = os.path.dirname(path_xlsx)
    if d:
        os.makedirs(d, exist_ok=True)
    with pd.ExcelWriter(path_xlsx, engine="openpyxl") as w:
        if len(bang_ma):
            bang_ma.to_excel(w, sheet_name="Tóm tắt theo mã", index=False)
        thong_ke(nk).to_excel(w, sheet_name="Thống kê", index=False)
        if len(ct):
            ct.to_excel(w, sheet_name="Chi tiết", index=False)
        nk.drop(columns=["id"]).to_excel(w, sheet_name="Nhật ký", index=False)
        trang_tri_excel(w, "Nhật ký")
        _to_mau_excel(w)
    return path_xlsx


def _to_mau_excel(w):
    """Định dạng số, ngày và tô xanh/đỏ cột % trong 2 sheet tóm tắt."""
    from openpyxl.styles import Font, PatternFill
    for ten in ("Tóm tắt theo mã", "Chi tiết"):
        if ten not in w.book.sheetnames:
            continue
        ws = w.book[ten]
        tieu_de = [o.value for o in ws[1]]
        for k, c in enumerate(tieu_de):
            for hang in ws.iter_rows(min_row=2, min_col=k + 1, max_col=k + 1):
                o = hang[0]
                if c in ("Thay đổi %", "Từ lần đầu %", "Lệnh giả định %") and isinstance(o.value, (int, float)):
                    o.number_format = '+0.0"%";-0.0"%"'
                    if o.value >= 1:
                        o.font, o.fill = Font(bold=True, color="1B5E20"), PatternFill("solid", fgColor="E8F5E9")
                    elif o.value <= -1:
                        o.font, o.fill = Font(bold=True, color="B71C1C"), PatternFill("solid", fgColor="FFEBEE")
                elif c in ("Giá lúc KN", "Giá hiện tại", "Cắt lỗ", "Mục tiêu"):
                    o.number_format = "#,##0.00"
                elif c in ("Ngày KN", "Lần đầu", "Giá tới ngày"):
                    o.number_format = "dd/mm/yyyy"
                elif c == "Đánh giá tạm" and isinstance(o.value, str):
                    if o.value.startswith("✔"):
                        o.font = Font(bold=True, color="1B5E20")
                    elif o.value.startswith("✘"):
                        o.font = Font(bold=True, color="B71C1C")


# ------------------------------------------------------------------ dùng trong ptcp.main()
def ghi_va_cham_ma(symbol, df_ngay, khuyen_nghi, gia, cat_lo, muc_tieu, ky_han, san="HOSE", ev=None, rr=None,
                   path=None):
    """
    Gọi từ main(): chấm lại khuyến nghị cũ của mã bằng df_ngay, rồi ghi khuyến nghị hôm nay (nếu đổi).
    Trả {"bang": các dòng của mã, "tk": thống kê của mã, "tk_tat_ca": thống kê toàn nhật ký, "path", "moi"}.
    """
    path = path or duong_dan()
    cap_nhat(None, {symbol: df_ngay}, [path], chi_ma=symbol)
    moi = ghi_khuyen_nghi(symbol, df_ngay.index[-1], gia, khuyen_nghi, cat_lo, muc_tieu, ky_han, san, ev, rr,
                          path=path)
    nk = doc(path)
    bang = nk[nk["ma"] == symbol]
    return {"bang": bang, "tk": thong_ke(bang), "tk_tat_ca": thong_ke(nk), "path": path, "moi": moi}
