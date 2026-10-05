# -*- coding: utf-8 -*-
"""Trang trí file Excel theo tông XANH LÁ (dùng chung cho ptcp & danh mục): tiêu đề xanh đậm chữ trắng,
dòng xen kẽ xanh nhạt, độ rộng cột theo nội dung, cố định hàng tiêu đề, nút lọc, tab màu xanh.
Ô đã được tô màu riêng (VD bảng 10 tiêu chí J5) được giữ nguyên màu."""
from .mau import XANH_DAM_HEX, XANH_HEX, XANH_NHAT_HEX


def _co_mau_rieng(o):
    try:
        return o.fill is not None and o.fill.fill_type == "solid" and \
            str(o.fill.fgColor.rgb) not in ("00000000", "FFFFFFFF", "None")
    except Exception:
        return False


def _sheet_tong_quan(wb, the, anh, tieu_de):
    from openpyxl.styles import Font, PatternFill
    from .so import vn_hoa
    ws = wb.create_sheet("Tong quan", 0)
    ws.sheet_properties.tabColor = XANH_DAM_HEX
    ws.cell(1, 1, vn_hoa(tieu_de)).font = Font(bold=True, size=16, color="FF" + XANH_DAM_HEX)
    dong = 3
    for nhan, gt, *_ in the or []:
        a, b = ws.cell(dong, 1, nhan), ws.cell(dong, 2, vn_hoa(str(gt)))
        a.font, a.fill = Font(bold=True, color="FFFFFFFF"), PatternFill("solid", fgColor="FF" + XANH_DAM_HEX)
        b.font, b.fill = Font(bold=True, size=12, color="FF" + XANH_DAM_HEX), PatternFill("solid", fgColor="FF" + XANH_NHAT_HEX)
        dong += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 40
    try:
        from openpyxl.drawing.image import Image as AnhXL
        import os
        for p in anh or []:
            if p and os.path.exists(p):
                a = AnhXL(p)
                ty = min(1.0, 950 / a.width)
                a.width, a.height = int(a.width * ty), int(a.height * ty)
                ws.add_image(a, f"A{dong + 1}")
                dong += int(a.height / 20) + 2
    except Exception:
        pass
    return ws


def trang_tri_excel(path, sheet_chinh=("Tom tat", "Danh muc", "Tong quan"), the=None, anh=None, tieu_de=""):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import ColorScaleRule
        from .so import la_cot_ty_le, vn_hoa
        wb = load_workbook(path)
    except Exception:
        return False
    nen_td = PatternFill("solid", fgColor="FF" + XANH_DAM_HEX)
    nen_xen = PatternFill("solid", fgColor="FF" + XANH_NHAT_HEX)
    vien = Border(bottom=Side(style="thin", color="FFC8E6C9"))
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = XANH_HEX if ws.title in sheet_chinh or ws.title.startswith("J") else "A5D6A7"
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        for o in ws[1]:
            o.fill, o.font = nen_td, Font(bold=True, color="FFFFFFFF")
            o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                o = ws.cell(r, c)
                o.border = vien
                if r % 2 == 1 and not _co_mau_rieng(o):
                    o.fill = nen_xen
                if isinstance(o.value, str):
                    o.value = vn_hoa(o.value)
                elif isinstance(o.value, float) and o.number_format == "General":
                    o.number_format = ("0.####" if la_cot_ty_le(ws.cell(1, c).value)
                                       else "#,##0.00" if abs(o.value) < 1e6 else "#,##0")
        for c in range(1, ws.max_column + 1):
            dai = max(len(str(ws.cell(r, c).value or "")) for r in range(1, min(ws.max_row, 300) + 1))
            ws.column_dimensions[get_column_letter(c)].width = max(8, min(dai + 2, 60))
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        if ws.title.startswith("J4") or "Bien dong" in ws.title or "Hieu suat" in ws.title:   # thang màu đỏ–xanh
            for c in range(2, ws.max_column + 1):
                if any(isinstance(ws.cell(r, c).value, (int, float)) for r in range(2, ws.max_row + 1)):
                    vung = f"{get_column_letter(c)}2:{get_column_letter(c)}{ws.max_row}"
                    ws.conditional_formatting.add(vung, ColorScaleRule(
                        start_type="min", start_color="FFEF9A9A", mid_type="num", mid_value=0, mid_color="FFFFFFFF",
                        end_type="max", end_color="FFA5D6A7"))
    if the or anh:
        _sheet_tong_quan(wb, the, anh, tieu_de)
    wb.save(path)
    return True
