# -*- coding: utf-8 -*-
"""Tóm tắt nhật ký theo mã: giá lúc khuyến nghị → hiện tại, làm mới & ghi file (không mạng)."""
import os

from openpyxl import load_workbook

from ptcp import nhat_ky, cau_hinh as cfg
from tests.conftest import tao_gia


def test_tom_tat_theo_ma_va_xuat(tmp_path):
    data = {"AAA": tao_gia(seed=3), "BBB": tao_gia(seed=5)}
    for ma, df in data.items():
        for lui, kn in ((40, "MUA"), (10, "CHƯA MUA")):
            d = df.iloc[:-lui]
            g = d.close.iloc[-1]
            nhat_ky.ghi_khuyen_nghi(ma, d.index[-1], g, kn, g * 0.93, g * 1.15, 63)
    kq = nhat_ky.xem(tai_gia=lambda ma, start: data[ma], chi_tiet=True)
    b = kq["theo_ma"]
    assert set(b["Mã"]) == {"AAA", "BBB"} and (b["Số KN"] == 2).all()
    for _, r in b.iterrows():
        gia_kn = data[r["Mã"]].close.iloc[-11]
        assert abs(r["Thay đổi %"] - (data[r["Mã"]].close.iloc[-1] / round(gia_kn, 2) - 1) * 100) < 1e-6
        assert r["Khuyến nghị gần nhất"] == "CHƯA MUA"
    assert b["Thay đổi %"].is_monotonic_decreasing
    # file nhật ký được ghi lại: dòng đang chờ có lãi/lỗ tạm tính
    nk = nhat_ky.doc()
    assert nk["ket_qua"].ne("").all()
    assert os.path.dirname(kq["file_excel"]) == os.path.dirname(os.path.abspath(cfg.FILE_NHAT_KY))
    assert load_workbook(kq["file_excel"]).sheetnames[0] == "Tóm tắt theo mã"


def test_danh_gia_tam():
    assert nhat_ky._danh_gia_tam("MUA", 5).startswith("✔")
    assert nhat_ky._danh_gia_tam("ĐỨNG NGOÀI", 5).startswith("✘")
    assert nhat_ky._danh_gia_tam("CHỜ", -5).startswith("✔")
    assert nhat_ky._danh_gia_tam("MUA", 0.3).startswith("–")
